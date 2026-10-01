"""Compare the SQL Server source structure with the PostgreSQL target structure.

Inputs are the column-metadata exports in ./inputs (one row per column with
schema_name, table_name, column_id, column_name, data_type, length_precision,
is_nullable, is_identity, is_primary_key, default_value).

Usage:  python compare_schemas.py  ->  writes schema_comparison.xlsx
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
SOURCE_FILE = HERE / "inputs" / "source_sqlserver_columns.xlsx"
TARGET_FILE = HERE / "inputs" / "target_postgres_columns.xlsx"
ROWCOUNT_FILE = HERE / "inputs" / "source_row_counts.xlsx"
OUTPUT_FILE = HERE / "schema_comparison.xlsx"

TABLE_MAP = {
    "date_table": "date_table",
    "District_Directors": "district_directors",
    "EmployeePaySummaryV2": "employee_pay_summary",
    "NetSuiteLocation_Mapping": "netsuite_location_mapping",
    "POS_ORDERDETAILS": "pos_order_details",
    "POS_ORDERPAIDOUTS": "pos_order_paid_outs",
    "POS_ORDERPAYMENTS": "pos_order_payments",
    "POS_ORDERS": "pos_orders",
    "Temp_DLH": "temp_dlh",
    "WeeklyCogs": "weekly_cogs",
    "WeeklyCogsProdNum": "weekly_cogs_prod_num",
    "GM_2026_UnPivot_New": "vena_gross_margin",
    "DLH_2026_UnPivot_New": "vena_labour_hours",
    "Sales_2026_UnPivot_New": "vena_sales",
    "Transactions_2026_UnPivot_v2": "vena_transactions",
}

# Source column names whose target name is not a plain snake_case conversion.
COLUMN_OVERRIDES = {
    "Endday": "end_day",
    "EndDay": "end_day",
    "SubTotal": "subtotal",
    "Discountamount": "discount_amount",
    "AcquitsionDate": "acquisition_date",
    "Day of_Week": "day_of_week",
    "PriceBackUP": "price_backup",
    "WeeklyCogs_ID": "id",
}

# Analyst notes for specific columns: (source_table, source_column or target_column) -> note
NOTES = {
    ("EmployeePaySummaryV2", "PayDate"): "datetime stored as varchar(20) in target - agree a text format (ISO 8601?) or change target to timestamp/date.",
    ("EmployeePaySummaryV2", "Retrieve_Date"): "datetime stored as varchar(30) in target - agree a text format or change target to timestamp.",
    ("EmployeePaySummaryV2", "AuthorizedManager"): "Not in target - confirm the column is intentionally dropped.",
    ("EmployeePaySummaryV2", "AuthorizedEmployee"): "Not in target - confirm the column is intentionally dropped.",
    ("WeeklyCogs", "createdate"): "Not in target by name - map to created_at? Otherwise the value is lost.",
    ("WeeklyCogs", "WeeklyCogs_ID"): "Identity value copied into id (bigint sequence); reset weekly_cogs_id_seq after load.",
    ("WeeklyCogs", "Start_Date"): "datetime -> date drops the time part; check source has no non-midnight times.",
    ("WeeklyCogs", "period"): "datetime -> date drops the time part; check source has no non-midnight times.",
    ("WeeklyCogsProdNum", "WeeklyCogsProdNum_ID"): "Identity column dropped; product_num becomes the primary key - source must have no NULL/duplicate product_num.",
    ("WeeklyCogsProdNum", "product_num"): "Target primary key - must be unique and non-null in source.",
    ("NetSuiteLocation_Mapping", "location_id_src"): "New column - define the source value (LocationID as text?).",
    ("NetSuiteLocation_Mapping", "LocationID"): "Source identity; target is a plain integer with no PK - consider making it the PK.",
}
GENERIC_NOTES = {
    "GUID": "varchar -> uuid: every value must be a valid UUID string; target is also NOT NULL.",
    "OrderTypeName": "varchar -> enum (USER-DEFINED): every distinct source value must exist in the enum; get enum labels.",
    "HasMods": "varchar flag -> boolean: map source values (e.g. Y/N, True/False, 1/0) explicitly.",
    "TaxExemption": "varchar flag -> boolean: map source values explicitly.",
    "Deposit": "varchar flag -> boolean: map source values explicitly.",
    "UsePennyRounding": "varchar flag -> boolean: map source values explicitly.",
    "IsRefund": "varchar flag -> boolean: map source values explicitly.",
    "StoreId": "Nullable in source, NOT NULL in target - check for NULLs.",
    "budget_year": "NOT NULL, no default, no source column - derive from the source table name (Sales_20xx_UnPivot_New -> 20xx); one source table per budget year.",
    "id": "Generated surrogate key (default) - not loaded from source.",
}
AUDIT_COLUMNS = {
    "created_at", "updated_at", "created_by", "updated_by",
    "created_timestamp", "updated_timestamp",
}

TYPE_FAMILY = {
    "int": "integer", "integer": "integer", "smallint": "smallint", "bigint": "bigint",
    "tinyint": "smallint",
    "varchar": "varchar", "nvarchar": "varchar", "character varying": "varchar",
    "char": "char", "nchar": "char", "character": "char",
    "numeric": "numeric", "decimal": "numeric",
    "float": "double", "double precision": "double", "real": "real",
    "date": "date",
    "smalldatetime": "timestamp", "datetime": "timestamp", "datetime2": "timestamp",
    "timestamp without time zone": "timestamp",
    "timestamp with time zone": "timestamptz",
    "uuid": "uuid", "boolean": "boolean", "bit": "boolean", "USER-DEFINED": "enum",
}
INT_RANK = {"smallint": 1, "integer": 2, "bigint": 3}

STATUS_FILL = {
    "Match": "C6EFCE",
    "Widened (safe)": "DDEBF7",
    "Renamed only": "DDEBF7",
    "Type change": "FFEB9C",
    "Length narrowed": "FFC7CE",
    "NULL -> NOT NULL": "FFC7CE",
    "Missing in target": "FFC7CE",
    "New in target": "EDEDED",
}


def load_rows(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if r and r[0]]


def snake(name: str) -> str:
    if name in COLUMN_OVERRIDES:
        return COLUMN_OVERRIDES[name]
    s = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", name)
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s)
    return s.replace(" ", "_").lower()


def is_nullable(value) -> bool:
    return str(value).strip().lower() in ("1", "true", "yes")


def parse_len(value) -> tuple[int, int] | int | None:
    if value in (None, ""):
        return None
    text = str(value)
    if "," in text:
        p, s = text.split(",")
        return int(p), int(s)
    return int(text)


def classify(src: dict, tgt: dict) -> tuple[list[str], str]:
    """Return (statuses, detail) for a mapped column pair."""
    statuses = []
    details = []
    sf = TYPE_FAMILY.get(src["data_type"], src["data_type"])
    tf = TYPE_FAMILY.get(tgt["data_type"], tgt["data_type"])
    sl, tl = parse_len(src["length_precision"]), parse_len(tgt["length_precision"])

    if sf == tf:
        if sf in ("varchar", "char") and sl and tl:
            if tl < sl:
                statuses.append("Length narrowed")
                details.append(f"{sl} -> {tl} characters: values longer than {tl} will fail")
            elif tl > sl:
                statuses.append("Widened (safe)")
        elif sf == "numeric" and sl and tl and sl != tl:
            (sp, ss), (tp, ts) = sl, tl
            if tp - ts < sp - ss or ts < ss:
                statuses.append("Length narrowed")
                details.append(f"numeric({sp},{ss}) -> numeric({tp},{ts})")
            else:
                statuses.append("Widened (safe)")
        elif sf == "timestamp" and src["data_type"] != tgt["data_type"]:
            statuses.append("Widened (safe)")
    elif sf in INT_RANK and tf in INT_RANK:
        statuses.append("Widened (safe)" if INT_RANK[tf] > INT_RANK[sf] else "Length narrowed")
    else:
        statuses.append("Type change")
        if sf == "double" and tf == "numeric":
            p, s = tl
            details.append(f"float -> numeric({p},{s}): rounds to {s} decimals; max {10 ** (p - s) - 1:,} before overflow")
        elif sf == "timestamp" and tf == "date":
            details.append("time of day is dropped")
        elif sf == "timestamp" and tf == "varchar":
            details.append(f"date/time stored as text (varchar({tl}))")

    if is_nullable(src["is_nullable"]) and not is_nullable(tgt["is_nullable"]):
        statuses.append("NULL -> NOT NULL")
        details.append("source allows NULL, target does not")

    if not statuses:
        statuses.append("Match" if snake(src["column_name"]) == src["column_name"] or src["column_name"].lower() == tgt["column_name"] else "Renamed only")
    return statuses, "; ".join(details)


def build_mapping(src_rows, tgt_rows):
    tgt_by_table: dict[str, dict[str, dict]] = {}
    for r in tgt_rows:
        tgt_by_table.setdefault(r["table_name"], {})[r["column_name"]] = r

    out = []
    for s_table, t_table in TABLE_MAP.items():
        s_cols = sorted((r for r in src_rows if r["table_name"] == s_table), key=lambda r: int(r["column_id"]))
        t_cols = tgt_by_table.get(t_table, {})
        used = set()
        for c in s_cols:
            want = snake(c["column_name"])
            match = next((t for k, t in t_cols.items() if k == c["column_name"] or k.lower() == want), None)
            note = NOTES.get((s_table, c["column_name"])) or GENERIC_NOTES.get(c["column_name"], "")
            if match is None:
                out.append(dict(src_table=s_table, tgt_table=t_table, src=c, tgt=None,
                                status="Missing in target", detail="", note=note))
                continue
            used.add(match["column_name"])
            statuses, detail = classify(c, match)
            out.append(dict(src_table=s_table, tgt_table=t_table, src=c, tgt=match,
                            status=" + ".join(statuses), detail=detail, note=note))
        for name, t in sorted(t_cols.items(), key=lambda kv: int(kv[1]["column_id"])):
            if name in used:
                continue
            note = NOTES.get((s_table, name)) or GENERIC_NOTES.get(name, "")
            if not note and name in AUDIT_COLUMNS:
                default = t["default_value"]
                has_default = default not in (None, "", "(null)")
                note = "Audit column" + (f" - filled by default {default}" if has_default else " - no default; populate or leave NULL")
                if not has_default and not is_nullable(t["is_nullable"]):
                    note += " (NOT NULL: must be populated)"
            out.append(dict(src_table=s_table, tgt_table=t_table, src=None, tgt=t,
                            status="New in target", detail="", note=note))
    return out


def write_workbook(mapping, row_counts):
    wb = openpyxl.Workbook()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    wrap = Alignment(wrap_text=True, vertical="top")

    # --- Column mapping sheet ---
    ws = wb.active
    ws.title = "Column Mapping"
    headers = ["Source table", "Target table", "Source column", "Source type", "Source len/prec",
               "Source nullable", "Target column", "Target type", "Target len/prec", "Target nullable",
               "Target default", "Status", "Detail", "Action / note"]
    ws.append(headers)
    for m in mapping:
        s, t = m["src"] or {}, m["tgt"] or {}
        default = t.get("default_value")
        ws.append([
            m["src_table"], m["tgt_table"],
            s.get("column_name", ""), s.get("data_type", ""), s.get("length_precision") or "",
            ("YES" if is_nullable(s["is_nullable"]) else "NO") if s else "",
            t.get("column_name", ""), t.get("data_type", ""), t.get("length_precision") or "",
            ("YES" if is_nullable(t["is_nullable"]) else "NO") if t else "",
            "" if default in (None, "(null)") else default,
            m["status"], m["detail"], m["note"],
        ])
    widths = [26, 26, 28, 14, 10, 9, 30, 26, 10, 9, 26, 30, 50, 70]
    _format_sheet(ws, headers, widths, font, bold, head_fill, wrap)
    status_col = headers.index("Status") + 1
    for row in ws.iter_rows(min_row=2):
        cell = row[status_col - 1]
        worst = _worst_status(cell.value)
        cell.fill = PatternFill("solid", fgColor=STATUS_FILL[worst])
    last_row = ws.max_row

    # --- Summary sheet (formulas over the mapping sheet) ---
    sm = wb.create_sheet("Summary", 0)
    sh = ["Source table", "Target table", "Source rows", "Source columns", "Mapped columns",
          "Exact / renamed / widened", "Type changes", "Narrowed length/precision",
          "NULL -> NOT NULL", "Missing in target", "New in target", "Like-for-like?"]
    sm.append(sh)
    rng = f"'Column Mapping'!$A$2:$A${last_row}"
    st = f"'Column Mapping'!$L$2:$L${last_row}"
    for i, (s_table, t_table) in enumerate(TABLE_MAP.items(), start=2):
        a = f"$A{i}"
        sm.append([
            s_table, t_table, row_counts.get(s_table.lower(), None),
            f'=COUNTIFS({rng},{a},\'Column Mapping\'!$C$2:$C${last_row},"<>")',
            f'=COUNTIFS({rng},{a},\'Column Mapping\'!$C$2:$C${last_row},"<>",\'Column Mapping\'!$G$2:$G${last_row},"<>")',
            f'=COUNTIFS({rng},{a},{st},"Match")+COUNTIFS({rng},{a},{st},"Renamed only")+COUNTIFS({rng},{a},{st},"Widened (safe)")',
            f'=COUNTIFS({rng},{a},{st},"*Type change*")',
            f'=COUNTIFS({rng},{a},{st},"*Length narrowed*")',
            f'=COUNTIFS({rng},{a},{st},"*NOT NULL*")',
            f'=COUNTIFS({rng},{a},{st},"Missing in target")',
            f'=COUNTIFS({rng},{a},{st},"New in target")',
            f'=IF(G{i}+H{i}+I{i}+J{i}+K{i}=0,"Yes",IF(G{i}+H{i}+I{i}+J{i}=0,"Yes + new columns","No"))',
        ])
    n = len(TABLE_MAP) + 1
    sm.append(["TOTAL", "", f"=SUM(C2:C{n})"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{n})" for c in range(4, 12)] + [""])
    _format_sheet(sm, sh, [30, 28, 14, 10, 10, 14, 10, 12, 12, 10, 10, 18], font, bold, head_fill, wrap)
    for c in sm[n + 1]:
        c.font = Font(name="Arial", size=10, bold=True)
    for r in range(2, n + 2):
        sm.cell(r, 3).number_format = "#,##0"
    sm.cell(n + 3, 1, "Source rows come from the user-supplied row-count export; all other counts are formulas over 'Column Mapping'.").font = Font(name="Arial", size=9, italic=True)
    sm.cell(n + 4, 1, "Status legend: green = match, blue = safe widening/rename, yellow = type conversion, red = data can fail to load or be lost, grey = new target column.").font = Font(name="Arial", size=9, italic=True)

    wb.calculation.fullCalcOnLoad = True  # Excel computes the Summary formulas on open
    wb.save(OUTPUT_FILE)


SEVERITY = ["Missing in target", "Length narrowed", "NULL -> NOT NULL", "Type change",
            "New in target", "Widened (safe)", "Renamed only", "Match"]


def _worst_status(value: str) -> str:
    for s in SEVERITY:
        if s in value:
            return s
    return "Match"


def _format_sheet(ws, headers, widths, font, bold, head_fill, wrap):
    for row in ws.iter_rows():
        for c in row:
            c.font = font
            c.alignment = wrap
    for i, h in enumerate(headers, start=1):
        c = ws.cell(1, i)
        c.font = bold
        c.fill = head_fill
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    src = load_rows(SOURCE_FILE)
    tgt = load_rows(TARGET_FILE)
    counts = {str(r["table_name"]).lower(): r["row_count"] for r in load_rows(ROWCOUNT_FILE)}
    mapping = build_mapping(src, tgt)
    write_workbook(mapping, counts)
    print(f"Wrote {OUTPUT_FILE} ({len(mapping)} mapping rows)")


if __name__ == "__main__":
    main()
